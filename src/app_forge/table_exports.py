"""Optional XLSX and PDF adapters for the framework-independent table model.

Install the corresponding extra (table-xlsx or table-pdf) to use an adapter.
The core package does not import either third-party library.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

__all__ = ["export_pdf", "export_xlsx"]


def _rows(model: Any, only_selected: bool) -> tuple[list[Any], list[dict[str, Any]]]:
    if not isinstance(only_selected, bool):
        raise TypeError("only_selected must be a boolean")
    columns = [column for column in model.columns if column.visible]
    rows = model.selected_rows() if only_selected else model.result_rows()
    return columns, rows


def _value(value: Any) -> Any:
    """Keep spreadsheet-native scalars; render compound values as JSON text."""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (list, tuple, dict)):
        try:
            return json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False)
        except (TypeError, ValueError):
            return str(value)
    return str(value)


def export_xlsx(model: Any, destination: str | Path, *, only_selected: bool = False) -> Path:
    """Write visible, filtered/sorted table rows to an XLSX workbook.

    Selected-only exports use the model's stable selection IDs. User-provided
    strings are explicitly stored as text so values beginning with '=' are
    not interpreted as spreadsheet formulas.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:
        raise RuntimeError("XLSX export requires the optional 'table-xlsx' extra (openpyxl)") from exc

    columns, rows = _rows(model, only_selected)
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Table"
    headers = [column.title or column.key for column in columns]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([_value(row.get(column.key)) for column in columns])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    if headers:
        sheet.auto_filter.ref = sheet.dimensions
        for index, header in enumerate(headers, start=1):
            letter = sheet.cell(row=1, column=index).column_letter
            longest = max([len(str(header)), *(len(str(row.get(columns[index - 1].key) or "")) for row in rows)])
            sheet.column_dimensions[letter].width = min(max(longest + 2, 10), 60)
    workbook.save(path)
    return path


def export_pdf(
    model: Any,
    destination: str | Path,
    *,
    title: str = "Table export",
    only_selected: bool = False,
) -> Path:
    """Write visible, filtered/sorted table rows to a paginated landscape PDF."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        from xml.sax.saxutils import escape
    except ImportError as exc:
        raise RuntimeError("PDF export requires the optional 'table-pdf' extra (reportlab)") from exc

    columns, rows = _rows(model, only_selected)
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    heading = styles["Title"]
    cell_style = ParagraphStyle(
        "TableExportCell", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=7, leading=9, alignment=TA_LEFT, spaceAfter=0,
    )
    header_style = ParagraphStyle(
        "TableExportHeader", parent=cell_style, fontName="Helvetica-Bold",
        textColor=colors.white,
    )
    data = [[Paragraph(escape(str(column.title or column.key)), header_style) for column in columns]]
    for row in rows:
        data.append([
            Paragraph(escape("" if row.get(column.key) is None else str(_value(row.get(column.key)))), cell_style)
            for column in columns
        ])
    page_size = landscape(A4)
    available_width = page_size[0] - 24 * mm
    widths = [available_width / len(columns)] * len(columns) if columns else []
    story = [Paragraph(escape(str(title)), heading), Spacer(1, 5 * mm)]
    if columns:
        table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#334155")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cbd5e1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(table)
    else:
        story.append(Paragraph("No visible columns", cell_style))
    SimpleDocTemplate(str(path), pagesize=page_size, rightMargin=12 * mm,
                      leftMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm).build(story)
    return path
